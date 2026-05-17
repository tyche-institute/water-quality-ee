"""
build_presentation.py — собирает submission/presentation.pptx из
готовых PNG-фигур (submission/figures/) и speaker notes (submission/speaker_notes.md).

Слайды (7 штук):
  1. Title
  2. Problem & Data           (figure: class_balance.png)
  3. Methodology              (text-only pipeline)
  4. Results — Model comparison (table from model_comparison.csv + roc_curves)
  5. Interpretation — SHAP & Calibration
  6. Limitations & Production (h2oatlas screenshot if present)
  7. Lessons Learned

Usage:
  .venv/bin/python scripts/build_presentation.py
"""

from __future__ import annotations

import json
import re
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parents[1]
SUB = ROOT / "submission"
FIG = SUB / "figures"
TBL = SUB / "tables"

BLUE   = RGBColor(0x0D, 0x6E, 0xFD)
DARK   = RGBColor(0x1F, 0x29, 0x37)
GRAY   = RGBColor(0x6C, 0x75, 0x7D)
LIGHT  = RGBColor(0xF1, 0xF5, 0xF9)
WHITE  = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xDC, 0x35, 0x45)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)


def _load_speaker_notes() -> dict[int, str]:
    """Parse submission/speaker_notes.md → {slide_index: notes_text}."""
    md = (SUB / "speaker_notes.md").read_text(encoding="utf-8")
    blocks: dict[int, str] = {}
    pattern = re.compile(r"^## Slide (\d+).*?$", re.MULTILINE)
    pieces = pattern.split(md)
    # pieces = [preamble, "1", body1, "2", body2, ...]
    for i in range(1, len(pieces), 2):
        idx = int(pieces[i])
        body = pieces[i + 1].strip()
        # remove following slide separators if any
        body = body.split("\n---")[0].strip()
        blocks[idx] = body
    return blocks


def _set_bg(slide, color: RGBColor) -> None:
    bg = slide.background
    fill = bg.fill
    fill.solid()
    fill.fore_color.rgb = color


def _add_text(slide, x, y, w, h, text, *, size=18, bold=False, color=DARK,
              align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Arial"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = Emu(0); tf.margin_right = Emu(0)
    tf.margin_top = Emu(0); tf.margin_bottom = Emu(0)
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.name = font
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    return tb


def _add_bullets(slide, x, y, w, h, items, *, size=18, color=DARK, font="Arial"):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.level = 0
        r = p.add_run()
        r.text = "•  " + line
        r.font.name = font
        r.font.size = Pt(size)
        r.font.color.rgb = color
        p.space_after = Pt(6)
    return tb


def _add_accent_bar(slide, top: bool = True) -> None:
    h = Inches(0.18)
    y = Inches(0) if top else SLIDE_H - h
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), y, SLIDE_W, h)
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()


def _add_title(slide, title: str, subtitle: str | None = None) -> None:
    _add_accent_bar(slide, top=True)
    _add_text(slide, Inches(0.6), Inches(0.35), SLIDE_W - Inches(1.2), Inches(0.6),
              title, size=28, bold=True, color=DARK)
    if subtitle:
        _add_text(slide, Inches(0.6), Inches(0.95), SLIDE_W - Inches(1.2), Inches(0.4),
                  subtitle, size=14, color=GRAY)


def _add_footer(slide, page: int, total: int) -> None:
    _add_text(slide, Inches(0.6), SLIDE_H - Inches(0.45),
              Inches(8), Inches(0.3),
              "Water Quality Compliance Estimator  ·  TalTech ITI8612  ·  h2oatlas.ee",
              size=9, color=GRAY)
    _add_text(slide, SLIDE_W - Inches(1.2), SLIDE_H - Inches(0.45),
              Inches(0.6), Inches(0.3), f"{page} / {total}",
              size=9, color=GRAY, align=PP_ALIGN.RIGHT)


def _attach_notes(slide, text: str) -> None:
    if not text:
        return
    notes = slide.notes_slide.notes_text_frame
    notes.clear()
    p = notes.paragraphs[0]
    r = p.add_run()
    r.text = text
    r.font.size = Pt(11)


def _new_slide(prs: Presentation):
    blank = prs.slide_layouts[6]   # blank layout
    s = prs.slides.add_slide(blank)
    _set_bg(s, WHITE)
    return s


# ── Slide builders ─────────────────────────────────────────────────────────────

def slide_1_title(prs, notes):
    s = _new_slide(prs)
    # Big top accent
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), SLIDE_W, Inches(2.2))
    bar.fill.solid(); bar.fill.fore_color.rgb = BLUE; bar.line.fill.background()

    _add_text(s, Inches(0.7), Inches(0.55), SLIDE_W - Inches(1.4), Inches(0.5),
              "TalTech ITI8612 · Machine Learning · Final Project · 2026",
              size=14, color=WHITE)
    _add_text(s, Inches(0.7), Inches(0.95), SLIDE_W - Inches(1.4), Inches(1.2),
              "Water Quality Compliance Estimator",
              size=40, bold=True, color=WHITE)
    _add_text(s, Inches(0.7), Inches(1.55), SLIDE_W - Inches(1.4), Inches(0.5),
              "A Probabilistic Risk Model for Estonian Open Water-Quality Data",
              size=18, color=WHITE)

    _add_text(s, Inches(0.7), Inches(2.7), SLIDE_W - Inches(1.4), Inches(0.4),
              "Anton Sokolov", size=22, bold=True, color=DARK)
    _add_text(s, Inches(0.7), Inches(3.15), SLIDE_W - Inches(1.4), Inches(0.4),
              "Supervisor: TalTech course staff", size=14, color=GRAY)

    _add_text(s, Inches(0.7), Inches(5.4), SLIDE_W - Inches(1.4), Inches(0.4),
              "Live application:", size=12, color=GRAY)
    _add_text(s, Inches(0.7), Inches(5.7), SLIDE_W - Inches(1.4), Inches(0.5),
              "https://h2oatlas.ee", size=22, bold=True, color=BLUE)
    _add_text(s, Inches(0.7), Inches(6.35), SLIDE_W - Inches(1.4), Inches(0.4),
              "Source: github.com/h2oatlas/water-quality-ee  ·  Data: vtiav.sm.ee",
              size=11, color=GRAY)
    _attach_notes(s, notes.get(1, ""))


def slide_2_problem_data(prs, notes, total):
    s = _new_slide(prs)
    _add_title(s, "Problem & Data",
               "Binary classification of laboratory probes against Estonian health norms")

    _add_bullets(s, Inches(0.6), Inches(1.6), Inches(7.0), Inches(4.5), [
        "Source: Terviseamet open data (vtiav.sm.ee), XML, 2021–2026",
        "4 domains: bathing, public water-supply, pools/SPA, drinking sources",
        "69 071 probes; ~70 engineered features",
        "Class imbalance: 86.9% compliant  vs  13.1% violation",
        "Priority metric: Recall on class 0 (violation)",
        "False negative ≫ False positive (E. coli, microbial risk)",
    ], size=18, color=DARK)

    s.shapes.add_picture(str(FIG / "class_balance.png"),
                         Inches(8.0), Inches(1.7),
                         width=Inches(4.8))
    _add_footer(s, 2, total)
    _attach_notes(s, notes.get(2, ""))


def slide_3_methodology(prs, notes, total):
    s = _new_slide(prs)
    _add_title(s, "Methodology",
               "Pipeline · feature engineering · 4 models · calibration · threshold tuning")

    # Pipeline as 6 boxes in a row
    steps = [
        ("XML\nload+parse", "data_loader.py"),
        ("Dedup\nlocation_key", "normalize_location()"),
        ("Feature\nengineering", "ratio + time + missing"),
        ("Split\n80/20 + TS-CV", "stratified + temporal"),
        ("Models\nLR · RF · GB · LightGBM", "class_weight balanced"),
        ("Calibrate\n+ threshold", "isotonic + PR tuning"),
    ]
    n = len(steps)
    pad = Inches(0.15)
    total_w = SLIDE_W - Inches(1.2)
    box_w = (total_w - pad * (n - 1)) / n
    box_h = Inches(1.4)
    y = Inches(2.0)
    for i, (label, sub) in enumerate(steps):
        x = Inches(0.6) + (box_w + pad) * i
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, box_w, box_h)
        shape.fill.solid(); shape.fill.fore_color.rgb = LIGHT
        shape.line.color.rgb = BLUE; shape.line.width = Pt(1.2)
        tf = shape.text_frame; tf.word_wrap = True
        tf.margin_left = Emu(45000); tf.margin_right = Emu(45000)
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label
        r.font.name = "Arial"; r.font.size = Pt(13); r.font.bold = True
        r.font.color.rgb = DARK
        p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
        r2 = p2.add_run(); r2.text = sub
        r2.font.size = Pt(9); r2.font.color.rgb = GRAY
        # Arrow (text)
        if i < n - 1:
            ax = x + box_w + Emu(10000)
            _add_text(s, ax, y + Inches(0.55), pad, Inches(0.3),
                      "→", size=18, color=GRAY, align=PP_ALIGN.CENTER)

    _add_text(s, Inches(0.6), Inches(4.0), SLIDE_W - Inches(1.2), Inches(0.5),
              "Key choices", size=20, bold=True, color=BLUE)
    _add_bullets(s, Inches(0.6), Inches(4.5), SLIDE_W - Inches(1.2), Inches(2.4), [
        "No SMOTE — keeps probability calibration interpretable",
        "Threshold optimised for Recall(viol.) at Precision ≥ 0.7",
        "County encoding is fit on train only (no leakage)",
        "Time-Series CV (n=5) confirms metrics survive temporal drift",
    ], size=15, color=DARK)
    _add_footer(s, 3, total)
    _attach_notes(s, notes.get(3, ""))


def slide_4_results(prs, notes, total):
    s = _new_slide(prs)
    _add_title(s, "Results — Model Comparison",
               "Hold-out test set (13 815 probes) · class-0 priority")

    df = pd.read_csv(TBL / "model_comparison.csv")
    df = df.set_index("model")
    headers = ["Model", "ROC-AUC", "Recall (viol.)", "Precision (viol.)", "F1 (viol.)"]
    rows = [[idx,
             f"{df.loc[idx, 'roc_auc']:.3f}",
             f"{df.loc[idx, 'recall_violation']:.3f}",
             f"{df.loc[idx, 'precision_violation']:.3f}",
             f"{df.loc[idx, 'f1_violation']:.3f}"] for idx in df.index]

    n_rows = len(rows) + 1
    n_cols = len(headers)
    tbl_x, tbl_y = Inches(0.6), Inches(1.7)
    tbl_w, tbl_h = Inches(7.0), Inches(2.6)
    table_shape = s.shapes.add_table(n_rows, n_cols, tbl_x, tbl_y, tbl_w, tbl_h)
    table = table_shape.table
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        cell.fill.solid(); cell.fill.fore_color.rgb = BLUE
        for p in cell.text_frame.paragraphs:
            for r in p.runs:
                r.font.bold = True; r.font.color.rgb = WHITE
                r.font.size = Pt(12); r.font.name = "Arial"
    for i, row in enumerate(rows, start=1):
        is_best = row[0] == "LightGBM"
        bg = LIGHT if is_best else WHITE
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = val
            cell.fill.solid(); cell.fill.fore_color.rgb = bg
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(12); r.font.name = "Arial"
                    r.font.bold = is_best
                    r.font.color.rgb = DARK

    _add_text(s, Inches(0.6), Inches(4.55), Inches(7.0), Inches(0.4),
              "LightGBM — best AUC + best F1", size=15, bold=True, color=ACCENT)
    _add_bullets(s, Inches(0.6), Inches(4.95), Inches(7.0), Inches(2.0), [
        "Recall 0.95 / Precision 0.90 at default threshold",
        "Threshold can slide along the PR curve at inference time",
        "GB reaches 0.96 recall at lower precision — alternative",
    ], size=14, color=DARK)

    s.shapes.add_picture(str(FIG / "roc_curves_4_models.png"),
                         Inches(7.85), Inches(1.7),
                         width=Inches(5.0))
    _add_footer(s, 4, total)
    _attach_notes(s, notes.get(4, ""))


def slide_5_interpretation(prs, notes, total):
    s = _new_slide(prs)
    _add_title(s, "Interpretation — SHAP & Calibration",
               "What drives the prediction · how trustworthy are the probabilities")

    s.shapes.add_picture(str(FIG / "shap_top10.png"),
                         Inches(0.4), Inches(1.6),
                         width=Inches(6.3))
    s.shapes.add_picture(str(FIG / "calibration_curve.png"),
                         Inches(7.0), Inches(1.6),
                         width=Inches(5.7))
    _add_bullets(s, Inches(0.4), Inches(6.0), SLIDE_W - Inches(0.8), Inches(1.2), [
        "Top-5: iron_missing · combined_chlorine · free_chlorine · colonies_37c_missing · ph_missing",
        "Missing-value indicators encode probe type → baseline risk",
        "All 4 models track the ideal diagonal closely — probabilities are trustworthy",
    ], size=12, color=DARK)
    _add_footer(s, 5, total)
    _attach_notes(s, notes.get(5, ""))


def slide_6_limitations_production(prs, notes, total):
    s = _new_slide(prs)
    _add_title(s, "Limitations & Production",
               "What the model can't do · how it ships to citizens")

    _add_text(s, Inches(0.6), Inches(1.6), Inches(6.5), Inches(0.5),
              "Three honest limitations", size=18, bold=True, color=BLUE)
    _add_bullets(s, Inches(0.6), Inches(2.05), Inches(6.5), Inches(3.5), [
        "Target = function of input → AUC > 0.99 reflects task structure",
        "Model does not measure unmeasured contaminants",
        "Phase-10 audit: 3.1% probes labelled violation with all params in norm",
        "Pool chlorine norms were wrong — fix lifted agreement +9 pp",
    ], size=14, color=DARK)

    _add_text(s, Inches(0.6), Inches(5.5), Inches(6.5), Inches(0.4),
              "Production:  h2oatlas.ee", size=18, bold=True, color=BLUE)
    _add_text(s, Inches(0.6), Inches(5.95), Inches(6.5), Inches(1.2),
              "2 196 locations · 2 layers (official + model risk) · "
              "3 languages · disclaimer that the service does not replace "
              "official assessment.", size=12, color=DARK)

    screenshot = FIG / "h2oatlas_screenshot.png"
    if screenshot.exists():
        s.shapes.add_picture(str(screenshot),
                             Inches(7.4), Inches(1.7),
                             width=Inches(5.4))
    else:
        # Placeholder card
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(7.4), Inches(1.7),
                                  Inches(5.4), Inches(4.5))
        card.fill.solid(); card.fill.fore_color.rgb = LIGHT
        card.line.color.rgb = BLUE; card.line.width = Pt(1.2)
        _add_text(s, Inches(7.4), Inches(3.6), Inches(5.4), Inches(0.6),
                  "[ insert h2oatlas.ee screenshot ]",
                  size=14, color=GRAY, align=PP_ALIGN.CENTER)
        _add_text(s, Inches(7.4), Inches(4.0), Inches(5.4), Inches(0.5),
                  "save as submission/figures/h2oatlas_screenshot.png",
                  size=10, color=GRAY, align=PP_ALIGN.CENTER)
    _add_footer(s, 6, total)
    _attach_notes(s, notes.get(6, ""))


def slide_7_lessons(prs, notes, total):
    s = _new_slide(prs)
    _add_title(s, "Lessons Learned",
               "Direct response to the supervisor's note: scope, data quality, imbalance")

    quote = ('"Data processing, handling imbalanced classes, and comparing '
             'multiple models may take more time than originally expected."')
    qbox = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                              Inches(0.6), Inches(1.6),
                              SLIDE_W - Inches(1.2), Inches(1.0))
    qbox.fill.solid(); qbox.fill.fore_color.rgb = LIGHT
    qbox.line.color.rgb = BLUE; qbox.line.width = Pt(1.2)
    tf = qbox.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = quote
    r.font.italic = True; r.font.size = Pt(14); r.font.color.rgb = DARK
    p2 = tf.add_paragraph(); p2.alignment = PP_ALIGN.CENTER
    r2 = p2.add_run(); r2.text = "— supervisor feedback"
    r2.font.size = Pt(10); r2.font.color.rgb = GRAY

    items = [
        ("1.  Progressive scope",
         "Started with 2 domains (supluskoha + veevärk). Expanded to 4 only "
         "after the pipeline stabilised."),
        ("2.  Data-quality audit (Phase 10)",
         "Added only after the first model surprised us — caught a real "
         "bug in pool chlorine norms (+9 pp agreement)."),
        ("3.  Imbalance handling",
         "class_weight='balanced' + threshold tuning, NOT SMOTE — keeps "
         "probabilities calibrated and interpretable on the public map."),
    ]
    y = Inches(3.0)
    for title, body in items:
        _add_text(s, Inches(0.6), y, SLIDE_W - Inches(1.2), Inches(0.45),
                  title, size=18, bold=True, color=BLUE)
        _add_text(s, Inches(0.6), y + Inches(0.45), SLIDE_W - Inches(1.2), Inches(0.7),
                  body, size=13, color=DARK)
        y += Inches(1.25)

    _add_text(s, Inches(0.6), SLIDE_H - Inches(0.85),
              SLIDE_W - Inches(1.2), Inches(0.4),
              "Thank you — questions welcome.",
              size=14, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
    _add_footer(s, 7, total)
    _attach_notes(s, notes.get(7, ""))


def main() -> None:
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    notes = _load_speaker_notes()
    total = 7
    slide_1_title(prs, notes)
    slide_2_problem_data(prs, notes, total)
    slide_3_methodology(prs, notes, total)
    slide_4_results(prs, notes, total)
    slide_5_interpretation(prs, notes, total)
    slide_6_limitations_production(prs, notes, total)
    slide_7_lessons(prs, notes, total)

    out = SUB / "presentation.pptx"
    prs.save(str(out))
    print(f"Wrote {out.relative_to(ROOT)}  ({out.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
